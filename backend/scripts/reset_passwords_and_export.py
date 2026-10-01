#!/usr/bin/env python3
"""
Reset user passwords and export credentials to Excel.
Rules:
1. Non-superadmin users: Password = <employee_code><random 4-digit number>
   - Hashed using PBKDF2-HMAC-SHA256 (via backend.app.auth.hash_password)
   - Exported to an Excel spreadsheet for user distribution.
2. Superadmin:
   - Password = Strong password with Name + Strong Number + @ and #
   - Hashed and updated in DB
   - STRICTLY EXCLUDED from the Excel spreadsheet.
   - Printed only in the terminal/console output.
"""

import argparse
import os
import secrets
import sys
from datetime import datetime

# Add parent directory to sys.path so backend.app imports work
CURRENT_DIR = os.path.dirname(os.path.abspath(__file__))
BACKEND_DIR = os.path.dirname(CURRENT_DIR)
if BACKEND_DIR not in sys.path:
    sys.path.insert(0, BACKEND_DIR)

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.auth import hash_password, verify_password
from app.database import SessionLocal
from app.models import Role, User


def generate_random_4digit() -> str:
    """Generate cryptographically secure 4-digit number (1000-9999)."""
    return str(secrets.randbelow(9000) + 1000)


def generate_superadmin_password() -> str:
    """Generate strong password with Name + strong number + @ and #."""
    rand_num = secrets.randbelow(900000) + 100000  # 6-digit strong number
    return f"Eagle#Admin@{rand_num}#"


def create_excel_report(rows: list, output_path: str):
    """Create a beautifully formatted Excel file with credentials."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Employee Credentials"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Title Banner
    ws.merge_cells("A1:H1")
    title_cell = ws["A1"]
    title_cell.value = "KPI Portal - Employee Login Credentials"
    title_cell.font = Font(name="Segoe UI", size=14, bold=True, color="1E3A8A")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    # Generated metadata
    ws.merge_cells("A2:H2")
    meta_cell = ws["A2"]
    meta_cell.value = f"Generated on: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} | Total Accounts: {len(rows)} (Super Admin Excluded)"
    meta_cell.font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
    meta_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Headers
    headers = [
        "S.No",
        "Employee Code",
        "Full Name",
        "Email Address",
        "System Role",
        "Department",
        "Designation",
        "Assigned Password",
    ]

    header_row_idx = 4
    ws.row_dimensions[header_row_idx].height = 28

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(
        left=thin_border_side,
        right=thin_border_side,
        top=thin_border_side,
        bottom=thin_border_side,
    )

    for col_idx, header_text in enumerate(headers, start=1):
        cell = ws.cell(row=header_row_idx, column=col_idx, value=header_text)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = cell_border

    # Data rows
    row_alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    pwd_fill = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
    pwd_font = Font(name="Consolas", size=11, bold=True, color="1E40AF")
    regular_font = Font(name="Segoe UI", size=10, color="0F172A")

    for i, item in enumerate(rows, start=1):
        row_idx = header_row_idx + i
        ws.row_dimensions[row_idx].height = 22
        is_even = (i % 2 == 0)
        current_fill = row_alt_fill if is_even else None

        values = [
            i,
            item["employee_code"],
            item["name"],
            item["email"],
            item["role"].capitalize(),
            item["department"],
            item["designation"],
            item["new_password"],
        ]

        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = cell_border

            # Styling specific columns
            if col_idx == 1:  # S.No
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = regular_font
                if current_fill:
                    cell.fill = current_fill
            elif col_idx == 2:  # Employee Code
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(name="Consolas", size=10, bold=True, color="334155")
                if current_fill:
                    cell.fill = current_fill
            elif col_idx == 8:  # Password
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = pwd_font
                cell.fill = pwd_fill
            else:
                align_h = "center" if col_idx == 5 else "left"
                cell.alignment = Alignment(horizontal=align_h, vertical="center")
                cell.font = regular_font
                if current_fill:
                    cell.fill = current_fill

    # Auto-adjust column widths
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            # Skip title & subtitle rows for width calculation
            if cell.row in (1, 2):
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Freeze panes below header
    ws.freeze_panes = f"A{header_row_idx + 1}"

    # Save
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)


def process_passwords(
    execute: bool = False,
    superadmin_custom_password: str | None = None,
    output_excel: str | None = None,
):
    if not output_excel:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        output_excel = os.path.join(BACKEND_DIR, f"employee_credentials_{timestamp}.xlsx")

    db = SessionLocal()
    try:
        users = db.query(User).order_by(User.id.asc()).all()
        print(f"\n=======================================================")
        print(f"  KPI Portal - Password Reset & Excel Generation")
        print(f"=======================================================")
        print(f"Mode: {'[LIVE COMMIT]' if execute else '[DRY-RUN (NO DB CHANGES)]'}")
        print(f"Total Users Found: {len(users)}")

        superadmin_users = []
        regular_user_rows = []

        # Determine superadmin password
        superadmin_pwd = superadmin_custom_password or generate_superadmin_password()

        for u in users:
            if u.role == Role.superadmin:
                superadmin_users.append((u, superadmin_pwd))
                if execute:
                    u.password_hash = hash_password(superadmin_pwd)
                continue

            # Regular user password: <employee_code><random 4-digit number>
            emp_code = (u.employee_no or f"EMP{u.id:04d}").strip()
            rand_digits = generate_random_4digit()
            new_pwd = f"{emp_code}{rand_digits}"

            # Retrieve department & designation safely
            dept_name = ""
            desig_name = ""
            if u.designation:
                desig_name = u.designation.name or ""
                if u.designation.department:
                    dept_name = u.designation.department.name or ""

            regular_user_rows.append({
                "id": u.id,
                "employee_code": emp_code,
                "name": u.name,
                "email": u.email,
                "role": u.role.value,
                "department": dept_name,
                "designation": desig_name,
                "new_password": new_pwd,
            })

            if execute:
                u.password_hash = hash_password(new_pwd)

        # Generate Excel spreadsheet
        print(f"\n[1/3] Generating Excel File (Excluding Super Admin)...")
        create_excel_report(regular_user_rows, output_excel)
        print(f"  -> Successfully generated: {output_excel}")
        print(f"  -> Total user credentials in Excel: {len(regular_user_rows)}")

        # Commit DB changes if execute
        print(f"\n[2/3] Database Update...")
        if execute:
            db.commit()
            print("  -> Successfully committed all new password hashes to the database!")
        else:
            db.rollback()
            print("  -> Dry run complete. Database was NOT modified.")

        # Superadmin Credentials Box (Console Only)
        print(f"\n[3/3] Super Admin Credentials (CONFIDENTIAL - NOT in Excel):")
        print("+" + "-" * 62 + "+")
        print("|  SUPER ADMIN CREDENTIALS (ONLY KNOWN BY YOU)                |")
        print("+" + "-" * 62 + "+")
        for sa_user, sa_pwd in superadmin_users:
            print(f"|  Account Name : {sa_user.name:<43} |")
            print(f"|  Email        : {sa_user.email:<43} |")
            print(f"|  Employee Code: {(sa_user.employee_no or 'N/A'):<43} |")
            print(f"|  Password     : {sa_pwd:<43} |")
        print("+" + "-" * 62 + "+")
        print("|  NOTE: Keep this password safe. It is NOT saved in Excel.   |")
        print("+" + "-" * 62 + "+\n")

        print("Sample of 5 employee passwords generated:")
        for r in regular_user_rows[:5]:
            print(f"  - [{r['employee_code']}] {r['name']} ({r['email']}) -> {r['new_password']}")

        return {
            "total_users": len(users),
            "regular_users": len(regular_user_rows),
            "superadmin_count": len(superadmin_users),
            "excel_path": output_excel,
            "superadmin_password": superadmin_pwd,
        }

    except Exception as e:
        db.rollback()
        print(f"\n[ERROR] Transaction failed: {e}", file=sys.stderr)
        raise
    finally:
        db.close()


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Reset passwords and export user credentials.")
    parser.add_argument(
        "--execute",
        action="store_true",
        help="Actually apply changes to the database. Without this, runs in dry-run mode.",
    )
    parser.add_argument(
        "--superadmin-password",
        type=str,
        default=None,
        help="Custom password for superadmin (must contain letters, numbers, @ and #).",
    )
    parser.add_argument(
        "--output-excel",
        type=str,
        default=None,
        help="Path for generated Excel file.",
    )

    args = parser.parse_args()
    process_passwords(
        execute=args.execute,
        superadmin_custom_password=args.superadmin_password,
        output_excel=args.output_excel,
    )
